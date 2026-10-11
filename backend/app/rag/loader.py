"""Step 1 of RAG ingestion: read raw files (txt, md, pdf, docx) into plain text."""
from dataclasses import dataclass
from pathlib import Path

SUPPORTED = {".txt", ".md", ".pdf", ".docx"}


@dataclass
class Document:
    source: str  # file name - kept so answers can cite where they came from
    text: str


def _read_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return "\n\n".join((page.extract_text() or "") for page in reader.pages)


def _read_docx(path: Path) -> str:
    from docx import Document as DocxDocument

    doc = DocxDocument(str(path))
    return "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())


def load_file(path: Path) -> Document:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md"}:
        text = path.read_text(encoding="utf-8", errors="ignore")
    elif suffix == ".pdf":
        text = _read_pdf(path)
    elif suffix == ".docx":
        text = _read_docx(path)
    else:
        raise ValueError(f"Unsupported file type: {path.suffix}")
    return Document(source=path.name, text=text)


def load_directory(directory: Path) -> list[Document]:
    """Load every supported file under `directory` (skips empty files and README.md)."""
    docs: list[Document] = []
    for path in sorted(directory.rglob("*")):
        if not path.is_file() or path.suffix.lower() not in SUPPORTED:
            continue
        if path.name.lower() == "readme.md":
            continue
        doc = load_file(path)
        if doc.text.strip():
            docs.append(doc)
    return docs
